from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
import re

import openpyxl
from openpyxl.utils import column_index_from_string


FINANCIAL = "\uae08\uc735\uae30\uad00"
PRIVATE_BOND = "\uc0ac\ubaa8\uc0ac\ucc44"
GENERAL_LOAN = "\uc77c\ubc18\ucc28\uc785\uae08"
SCHEDULE_METHOD = "\uc0c1\ud658\uc2a4\ucf00\uc904"
BULLET_METHOD = "\ub9cc\uae30\uc77c\uc2dc\uc0c1\ud658"


def _safe_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _safe_amount(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return 0.0
    multiplier = 1.0
    if "\uc5b5" in text:
        multiplier = 100_000_000.0
    elif "\ub9cc" in text:
        multiplier = 10_000.0
    text = re.sub(r"[,%\uc6d0\uc5b5\ub9cc\s]", "", text)
    if not text:
        return 0.0
    try:
        return float(text) * multiplier
    except (TypeError, ValueError):
        return 0.0


def _safe_rate(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("%", "").replace(",", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except (TypeError, ValueError):
        return 0.0


def _safe_int(value) -> int | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return int(value)
    match = re.search(r"-?\d+", str(value).replace(",", ""))
    if not match:
        return None
    try:
        return int(match.group(0))
    except ValueError:
        return None


def _date_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        text = str(int(value))
    else:
        text = str(value or "").strip()
    if not text:
        return ""
    compact = re.sub(r"\D", "", text)
    if len(compact) == 8:
        return f"{compact[:4]}-{compact[4:6]}-{compact[6:8]}"
    if len(compact) == 6:
        return f"20{compact[:2]}-{compact[2:4]}-{compact[4:6]}"
    match = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", text)
    if match:
        return f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    match = re.search(r"(?<!\d)(\d{2})\D+(\d{1,2})\D+(\d{1,2})(?!\d)", text)
    if match:
        return f"20{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    return text


def _period_from_date(value: str) -> str:
    text = _date_str(value)
    match = re.search(r"(\d{4})-(\d{2})", text)
    return f"{match.group(1)}-{match.group(2)}" if match else ""


def _normalize_category_value(value: str) -> str:
    compact = str(value or "").strip().replace(" ", "")
    if not compact:
        return ""
    if "\uae08\uc735\uae30\uad00" in compact:
        return FINANCIAL
    if "\uc0ac\ubaa8\uc0ac\ucc44" in compact or compact == "\uc0ac\ucc44":
        return PRIVATE_BOND
    if "\uc77c\ubc18\ucc28\uc785" in compact:
        return GENERAL_LOAN
    return str(value or "").strip()


def _normalize_category(investor_name: str, contract_name: str, contract_kind: str, collateral_type: str, investor_type: str = "") -> str:
    explicit_category = _normalize_category_value(investor_type)
    if explicit_category in (FINANCIAL, PRIVATE_BOND, GENERAL_LOAN):
        return explicit_category

    text = f"{investor_name} {contract_name} {contract_kind} {collateral_type}".replace(" ", "")
    if "\uc0ac\ubaa8\uc0ac\ucc44" in text or "\uc0ac\ucc44" in text:
        return PRIVATE_BOND
    if "\uc77c\ubc18\ucc28\uc785" in text or ("\uc77c\ubc18" in text and "\ucc28\uc785" in text):
        return GENERAL_LOAN
    financial_keywords = (
        "\uc740\ud589",
        "\uc800\ucd95",
        "\uce90\ud53c\ud0c8",
        "\uce74\ub4dc",
        "\uae08\uc735",
        "\uc218\ud611",
        "\uc2e0\ud611",
        "\uc6b0\ub9ac",
        "\uad6d\ubbfc",
        "\uc2e0\ud55c",
        "\ud558\ub098",
        "\uae30\uc5c5",
    )
    if any(keyword in investor_name for keyword in financial_keywords):
        return FINANCIAL
    return GENERAL_LOAN


def _repayment_method(category: str) -> str:
    if category == FINANCIAL:
        return SCHEDULE_METHOD
    return BULLET_METHOD


def _normalize_repayment_method(value: str, category: str) -> str:
    text = _safe_text(value).replace(" ", "")
    if not text:
        return _repayment_method(category)
    if text == "\ub9cc\uae30\uc77c\uc2dc(\uc775\uc6d4)":
        return "\ub9cc\uae30\uc77c\uc2dc"
    return _safe_text(value)


def parse_borrowing_contract_list(filepath: str, period: str = "") -> dict:
    workbook = openpyxl.load_workbook(filepath, data_only=True)
    worksheet = workbook.active
    cols = {
        "key": column_index_from_string("A"),
        "investor_name": column_index_from_string("B"),
        "contract_name": column_index_from_string("C"),
        "contract_kind": column_index_from_string("D"),
        "contract_date": column_index_from_string("E"),
        "maturity_date": column_index_from_string("F"),
        "repayment_period": column_index_from_string("G"),
        "collateral_type": column_index_from_string("H"),
        "interest_rate": column_index_from_string("I"),
        "prepayment_fee_rate": column_index_from_string("J"),
        "agreement_day": column_index_from_string("K"),
        "borrow_amount": column_index_from_string("L"),
        "borrow_balance": column_index_from_string("M"),
        "next_payment_date": column_index_from_string("N"),
        "next_payment_amount": column_index_from_string("O"),
        "repayment_bank": column_index_from_string("P"),
        "investor_type": column_index_from_string("T"),
        "repayment_method": column_index_from_string("U"),
    }

    rows: list[dict] = []
    for row_idx in range(2, worksheet.max_row + 1):
        key_value = _safe_text(worksheet.cell(row=row_idx, column=cols["key"]).value)
        investor_name = _safe_text(worksheet.cell(row=row_idx, column=cols["investor_name"]).value)
        contract_name = _safe_text(worksheet.cell(row=row_idx, column=cols["contract_name"]).value)
        contract_kind = _safe_text(worksheet.cell(row=row_idx, column=cols["contract_kind"]).value)
        contract_date = _date_str(worksheet.cell(row=row_idx, column=cols["contract_date"]).value)
        maturity_date = _date_str(worksheet.cell(row=row_idx, column=cols["maturity_date"]).value)
        collateral_type = _safe_text(worksheet.cell(row=row_idx, column=cols["collateral_type"]).value)
        investor_type = _safe_text(worksheet.cell(row=row_idx, column=cols["investor_type"]).value)
        borrow_amount = _safe_amount(worksheet.cell(row=row_idx, column=cols["borrow_amount"]).value)
        borrow_balance = _safe_amount(worksheet.cell(row=row_idx, column=cols["borrow_balance"]).value)

        if borrow_amount == 0 and borrow_balance == 0:
            continue
        if not investor_name:
            continue

        category = _normalize_category(investor_name, contract_name, contract_kind, collateral_type, investor_type)
        row_period = period or _period_from_date(contract_date)
        interest_rate = _safe_rate(worksheet.cell(row=row_idx, column=cols["interest_rate"]).value)
        monthly_interest = (borrow_balance or borrow_amount) * interest_rate / 100 / 12 if interest_rate else 0.0
        agreement_day = _safe_int(worksheet.cell(row=row_idx, column=cols["agreement_day"]).value)
        rows.append({
            "period": row_period,
            "source_row": row_idx,
            "key": key_value,
            "category": category,
            "investor_type": investor_type or category,
            "investor_name": investor_name,
            "borrower": investor_name,
            "contract_name": contract_name,
            "round_label": contract_name,
            "contract_kind": contract_kind,
            "contract_date": contract_date,
            "maturity_date": maturity_date,
            "repayment_period": _safe_text(worksheet.cell(row=row_idx, column=cols["repayment_period"]).value),
            "repayment_period_months": _safe_int(worksheet.cell(row=row_idx, column=cols["repayment_period"]).value),
            "collateral_type": collateral_type,
            "interest_rate": interest_rate,
            "prepayment_fee_rate": _safe_rate(worksheet.cell(row=row_idx, column=cols["prepayment_fee_rate"]).value),
            "agreement_day": agreement_day,
            "borrow_amount": borrow_amount,
            "borrow_balance": borrow_balance,
            "next_payment_date": _date_str(worksheet.cell(row=row_idx, column=cols["next_payment_date"]).value),
            "next_payment_amount": _safe_amount(worksheet.cell(row=row_idx, column=cols["next_payment_amount"]).value),
            "repayment_method": _normalize_repayment_method(worksheet.cell(row=row_idx, column=cols["repayment_method"]).value, category),
            "repayment_agency": _safe_text(worksheet.cell(row=row_idx, column=cols["repayment_bank"]).value) or investor_name,
            "repayment_institution": _safe_text(worksheet.cell(row=row_idx, column=cols["repayment_bank"]).value) or investor_name,
            "monthly_interest_estimate": monthly_interest,
        })

    periods = sorted({
        str(row.get("period") or "").strip()
        for row in rows
        if str(row.get("period") or "").strip()
    })
    by_category: dict[str, dict] = {}
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["category"]].append(row)
    for category in (FINANCIAL, PRIVATE_BOND, GENERAL_LOAN):
        items = grouped.get(category, [])
        by_category[category] = {
            "rows": len(items),
            "borrow_amount": sum(float(row.get("borrow_amount") or 0) for row in items),
            "borrow_balance": sum(float(row.get("borrow_balance") or 0) for row in items),
            "monthly_interest_estimate": sum(float(row.get("monthly_interest_estimate") or 0) for row in items),
        }

    return {
        "period": period or (periods[-1] if periods else ""),
        "periods": periods,
        "rows": rows,
        "totals": {
            "rows": len(rows),
            "borrow_amount": sum(float(row.get("borrow_amount") or 0) for row in rows),
            "borrow_balance": sum(float(row.get("borrow_balance") or 0) for row in rows),
            "monthly_interest_estimate": sum(float(row.get("monthly_interest_estimate") or 0) for row in rows),
        },
        "by_category": by_category,
        "source": "borrowing_contracts",
    }
