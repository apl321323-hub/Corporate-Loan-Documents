"""
Payment statement parser.

Rules:
- K column accounting flag must be '+'.
- N column principal payment is aggregated as principal repayment.
- Q column interest payment is aggregated as interest repayment.
- E column payment date determines the period.
- B/H/AG/BC/BN columns are preserved for daily-close excellent lender matching.
- AK column loan date is preserved for daily-close usage-period buckets.
"""

import os
import re
from collections import defaultdict
from datetime import datetime, date

import openpyxl

COL_B = 2     # contract number
COL_E = 5     # payment date
COL_H = 8     # current product
COL_K = 11    # accounting flag
COL_N = 14    # principal payment
COL_Q = 17    # interest payment
COL_AG = 33   # collateral kind
COL_AK = 37   # loan date
COL_BC = 55   # NICE score
COL_BN = 66   # K score

KEY_PRINCIPAL = '\uc6d0\uae08\ud68c\uc218\uc561'
KEY_INTEREST = '\uc774\uc790\ud68c\uc218\uc561'
GROUP_CREDIT = '\uc2e0\uc6a9'
GROUP_COLLATERAL = '\ub2f4\ubcf4'


def _date_str(val) -> str | None:
    """Convert an Excel/date-like value to YYYY-MM-DD."""
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.strftime('%Y-%m-%d')
    s = str(val).strip()
    m = re.search(r'(\d{4})\D+(\d{1,2})\D+(\d{1,2})', s)
    if m:
        month = int(m.group(2))
        day = int(m.group(3))
        if 1 <= month <= 12 and 1 <= day <= 31:
            return f"{m.group(1)}-{month:02d}-{day:02d}"
    m = re.search(r'(?<!\d)(\d{2})\D+(\d{1,2})\D+(\d{1,2})(?!\d)', s)
    if m:
        month = int(m.group(2))
        day = int(m.group(3))
        if 1 <= month <= 12 and 1 <= day <= 31:
            return f"20{m.group(1)}-{month:02d}-{day:02d}"
    return None


def _ym(val) -> str | None:
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.strftime('%Y-%m')
    ds = _date_str(val)
    if ds:
        return ds[:7]
    s = str(val).strip()
    if len(s) >= 7:
        return s[:7]
    return None


def _date_obj(val) -> date | None:
    s = _date_str(val)
    if not s:
        return None
    try:
        return datetime.strptime(s, '%Y-%m-%d').date()
    except ValueError:
        return None


def _safe_amt(v) -> float:
    if v is None:
        return 0.0
    try:
        return float(v)
    except (ValueError, TypeError):
        text = str(v).replace(',', '').strip()
        try:
            return float(text)
        except (ValueError, TypeError):
            return 0.0


def _safe_score(v):
    if v is None:
        return None
    text = str(v).replace(',', '').strip()
    if not text or text in {'-', '―'}:
        return None
    try:
        return float(text)
    except (ValueError, TypeError):
        return None


def _cell_text(ws, row: int, col: int) -> str:
    if not col or col > ws.max_column:
        return ''
    value = ws.cell(row=row, column=col).value
    return str(value).strip() if value is not None else ''


def _norm_header(val) -> str:
    return re.sub(r'\s+', '', str(val or '')).lower()


def _find_col(ws, keywords, reserved: set[int] | None = None) -> int | None:
    reserved = reserved or set()
    for col in range(1, ws.max_column + 1):
        if col in reserved:
            continue
        label = _norm_header(ws.cell(row=1, column=col).value)
        if not label:
            continue
        if any(k.lower() in label for k in keywords):
            return col
    return None


def _resolve_group(raw_group, product_name: str, prod_to_group: dict[str, str]) -> str:
    text = str(raw_group or '').strip()
    if text in prod_to_group:
        return prod_to_group[text]
    if product_name in prod_to_group:
        return prod_to_group[product_name]
    if GROUP_CREDIT in text:
        return GROUP_CREDIT
    if GROUP_COLLATERAL in text:
        return GROUP_COLLATERAL
    return ''


def parse_payment(filepath: str, product_groups: list | None = None) -> dict:
    wb = openpyxl.load_workbook(filepath, data_only=True)
    source_file = os.path.basename(filepath)

    ws = None
    for sname in wb.sheetnames:
        if sname.strip() == '\uc785\uae08\uba85\uc138':
            ws = wb[sname]
            break
    if ws is None:
        for sname in wb.sheetnames:
            if '\uc785\uae08\uba85\uc138' in sname:
                ws = wb[sname]
                break
    if ws is None:
        ws = wb.active

    prod_to_group: dict[str, str] = {}
    for group in (product_groups or []):
        gname = str(group.get('name') or '').strip()
        items = group.get('products', group.get('items', []))
        for item in (items or []):
            pname = str(item or '').strip()
            if pname:
                prod_to_group[pname] = gname

    reserved_cols = {COL_B, COL_E, COL_H, COL_K, COL_N, COL_Q, COL_AG, COL_AK, COL_BC, COL_BN}
    contract_no_col = _find_col(
        ws,
        (
            '계약번호',
            '계약no',
            '계약번호',
            '대출번호',
            '채권번호',
            '관리번호',
        ),
        reserved_cols,
    )
    group_col = _find_col(
        ws,
        (
            '\uc0c1\ud488\uad6c\ubd84',
            '\ub2f4\ubcf4\uad6c\ubd84',
            '\uc2e0\uc6a9\ub2f4\ubcf4',
            '\uad6c\ubd84',
        ),
        reserved_cols,
    )
    product_col = _find_col(
        ws,
        (
            '\uc0c1\ud488\uba85',
            '\ub300\ucd9c\uc0c1\ud488',
            '\uc0c1\ud488',
        ),
        reserved_cols | ({group_col} if group_col else set()),
    )

    raw: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    raw_group: dict[str, dict[str, dict[str, float]]] = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    raw_product: dict[str, dict[str, dict[str, float]]] = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    raw_rows: list[dict] = []

    for row in range(2, ws.max_row + 1):
        k = ws.cell(row=row, column=COL_K).value
        if str(k or '').strip() != '+':
            continue

        e = ws.cell(row=row, column=COL_E).value
        ym = _ym(e)
        if not ym:
            continue

        n = ws.cell(row=row, column=COL_N).value
        q = ws.cell(row=row, column=COL_Q).value
        ak = ws.cell(row=row, column=COL_AK).value
        principal = _safe_amt(n)
        interest = _safe_amt(q)
        raw[ym][KEY_PRINCIPAL] += principal
        raw[ym][KEY_INTEREST] += interest

        contract_no_raw = _cell_text(ws, row, COL_B)
        if not contract_no_raw and contract_no_col:
            contract_no_raw = _cell_text(ws, row, contract_no_col)

        product_name = _cell_text(ws, row, COL_H)
        if not product_name and product_col:
            product_name = _cell_text(ws, row, product_col)

        group_raw = ''
        if group_col:
            group_val = ws.cell(row=row, column=group_col).value
            group_raw = str(group_val).strip() if group_val is not None else ''

        group_name = _resolve_group(group_raw, product_name, prod_to_group)
        if group_name:
            raw_group[group_name][ym][KEY_PRINCIPAL] += principal
            raw_group[group_name][ym][KEY_INTEREST] += interest
        if product_name:
            raw_product[product_name][ym][KEY_PRINCIPAL] += principal
            raw_product[product_name][ym][KEY_INTEREST] += interest
        pay_dt = _date_obj(e)
        loan_dt = _date_obj(ak)
        use_days = (pay_dt - loan_dt).days if pay_dt and loan_dt else None

        if principal:
            raw_rows.append({
                'period': ym,
                'date': _date_str(e) or '',
                'loan_date': _date_str(ak) or '',
                'contract_no': contract_no_raw,
                'product': product_name,
                'group': group_name,
                'accounting': str(k or '').strip(),
                'collateral_kind': _cell_text(ws, row, COL_AG),
                'nice_score': _safe_score(ws.cell(row=row, column=COL_BC).value) if COL_BC <= ws.max_column else None,
                'k_score': _safe_score(ws.cell(row=row, column=COL_BN).value) if COL_BN <= ws.max_column else None,
                'principal': principal / 1_000_000,
                'principal_won': principal,
                'use_days': use_days,
            })

    wb.close()

    periods = sorted(raw.keys())
    section1 = {KEY_PRINCIPAL: {}, KEY_INTEREST: {}}
    for ym in periods:
        d = raw[ym]
        section1[KEY_PRINCIPAL][ym] = d[KEY_PRINCIPAL] / 1_000_000
        section1[KEY_INTEREST][ym] = d[KEY_INTEREST] / 1_000_000

    def build_section(bucket: dict) -> dict:
        section: dict = {}
        for name in sorted(bucket.keys()):
            section[name] = {KEY_PRINCIPAL: {}, KEY_INTEREST: {}}
            for ym in periods:
                d = bucket[name].get(ym, {})
                section[name][KEY_PRINCIPAL][ym] = d.get(KEY_PRINCIPAL, 0.0) / 1_000_000
                section[name][KEY_INTEREST][ym] = d.get(KEY_INTEREST, 0.0) / 1_000_000
        return section

    return {
        'periods': periods,
        'section1': section1,
        'section3': build_section(raw_group),
        'section_product': build_section(raw_product),
        'raw_rows': raw_rows,
        'group_column': group_col,
        'product_column': product_col,
        'contract_no_column': contract_no_col,
        'source_file': source_file,
    }


if __name__ == '__main__':
    import json
    import sys

    fpath = sys.argv[1]
    print(json.dumps(parse_payment(fpath), ensure_ascii=False, indent=2))
