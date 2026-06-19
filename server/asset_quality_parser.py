"""
자산건전성 엑셀 파서
파일: 기업여신자료 - 자산건정성.xlsx  (시트: '자산건전성 (APL기준)')

구조:
  Row 1  : ■ 자산건전성 (헤더)
  Row 2  : (단위: 백만원)
  Row 3  : 구분 | 2016-12 | 2017-01 | ...  (기간 헤더)
  Row 4~15 : 섹션1 금액 (무연체~융잔합계)
  Row 16 : 검수행 (skip)
  Row 18~36: 섹션2 비율(%) (구분(%)~31일이상충당금)
  Row 38~75: 섹션3 담보구분별 (신용/보증/담보)
  Row 78~397: 섹션4 상품별 (29개 상품 × 9행)
"""

from datetime import datetime
import openpyxl


# ── 섹션 정의 ─────────────────────────────────────────────
SECTION1_ROWS = {
    4:  '무연체',
    5:  '1~30',
    6:  '1~10',
    7:  '11~30',
    8:  '31~60',
    9:  '61~90',
    10: '91~120',
    11: '121~150',
    12: '151~180',
    13: '181~',
    14: '연체합계',
    15: '융잔합계',
}

SECTION2_ROWS = {
    19: '무연체',
    20: '1~30',
    21: '1~10',
    22: '11~30',
    23: '31~60',
    24: '61~90',
    25: '91~120',
    26: '121~150',
    27: '151~180',
    28: '181~',
    29: '합계(%)',
    30: '대손충당금',
    31: '30일이상연체율',
    32: '연체합계',
    33: '연체율(%)',
    35: '31일이상연체합계',
    36: '31일이상충당금',
}

# 섹션3: 신용/보증/담보  각각 8행 + 연체율 행
SECTION3_GROUPS = {
    '신용': {'start': 38, 'rows': {
        40: '무연체', 41: '1~30', 42: '31~60', 43: '61~90',
        44: '91~180', 45: '181~', 46: '연체합계', 47: '융잔합계', 49: '연체율(%)',
    }},
    '보증': {'start': 51, 'rows': {
        53: '무연체', 54: '1~30', 55: '31~60', 56: '61~90',
        57: '91~180', 58: '181~', 59: '연체합계', 60: '융잔합계', 62: '연체율(%)',
    }},
    '담보': {'start': 64, 'rows': {
        66: '무연체', 67: '1~30', 68: '31~60', 69: '61~90',
        70: '91~180', 71: '181~', 72: '연체합계', 73: '융잔합계', 75: '연체율(%)',
    }},
}

# 섹션4: 상품별 (각 상품 블록 시작행)
SECTION4_PRODUCTS = [
    ('N론',        80),
    ('V론',        91),
    ('담보론',      102),
    ('담보론(지분대출)', 113),
    ('레이디론',    124),
    ('스타론',      135),
    ('스타스위치론', 146),
    ('우량론',      157),
    ('큐브론',      168),
    ('테일론',      179),
    ('토탈론',      190),
    ('프리론',      201),
    ('프리미엄론',   212),
    ('전월세론',    223),
    ('플러스론',    234),
    ('토마토론',    245),
    ('토마토N론',   256),
    ('토마토토탈론', 267),
    ('T플러스론',   278),
    ('OP론',       289),
    ('충성론',      300),
    ('오투론',      311),
    ('오투N론',     322),
    ('토마토토탈론플러스', 333),
    ('다이렉트론(A)', 344),
    ('다이렉트론(W)', 355),
    ('N론(하이브리드)', 366),
    ('기타',        377),
    ('기타N',       388),
]

PRODUCT_ROW_OFFSETS = {
    1: '무연체',
    2: '1~30',
    3: '31~60',
    4: '61~90',
    5: '91~120',
    6: '121~150',
    7: '151~180',
    8: '181~',
    9: '융잔합계',
}


def _safe_num(v):
    """엑셀 값을 숫자로 변환 (수식/문자열/None → None)"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    return None


def parse_asset_quality_excel(filepath: str) -> dict:
    """
    자산건전성 엑셀을 파싱하여 JSON 직렬화 가능한 dict 반환.

    반환 구조:
    {
      "periods": ["2016.12", "2017.01", ...],          # 103개
      "section1": { "무연체": [v1, v2, ...], ... },    # 금액 (백만원)
      "section2": { "무연체": [v1, ...], ... },         # 비율/금액 혼합
      "section3": {
        "신용": { "무연체": [...], ... },
        "보증": { ... },
        "담보": { ... }
      },
      "section4": {
        "N론": { "무연체": [...], ... },
        ...
      }
    }
    """
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    # ── 기간 헤더 파싱 (row 3, col 2~) ────────────────────
    periods = []
    col_indices = []  # 1-based column indices
    for c in range(2, ws.max_column + 1):
        v = ws.cell(3, c).value
        if v is None:
            continue
        if isinstance(v, datetime):
            periods.append(v.strftime('%Y.%m'))
        else:
            s = str(v).strip()
            if s:
                periods.append(s)
        col_indices.append(c)

    n = len(col_indices)

    def read_row(row_num):
        """지정 행의 모든 기간 값을 리스트로 반환"""
        return [_safe_num(ws.cell(row_num, col_indices[i]).value) for i in range(n)]

    # ── 섹션1 ─────────────────────────────────────────────
    section1 = {}
    for r, label in SECTION1_ROWS.items():
        section1[label] = read_row(r)

    # ── 섹션2 ─────────────────────────────────────────────
    section2 = {}
    for r, label in SECTION2_ROWS.items():
        section2[label] = read_row(r)

    # ── 섹션3 담보구분별 ──────────────────────────────────
    section3 = {}
    for grp_name, cfg in SECTION3_GROUPS.items():
        grp_data = {}
        for r, label in cfg['rows'].items():
            grp_data[label] = read_row(r)
        section3[grp_name] = grp_data

    # ── 섹션4 상품별 ──────────────────────────────────────
    section4 = {}
    for product_name, start_row in SECTION4_PRODUCTS:
        prod_data = {}
        for offset, label in PRODUCT_ROW_OFFSETS.items():
            prod_data[label] = read_row(start_row + offset - 1)
        section4[product_name] = prod_data

    return {
        "periods": periods,
        "section1": section1,
        "section2": section2,
        "section3": section3,
        "section4": section4,
        "products": [p for p, _ in SECTION4_PRODUCTS],
    }
