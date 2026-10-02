from __future__ import annotations

import datetime as _dt
import re
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter


def _safe_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _repair_mojibake(value: Any) -> str:
    text = _safe_text(value)
    if not text:
        return ""
    try:
        repaired = text.encode("latin-1").decode("utf-8")
        return repaired if repaired else text
    except UnicodeError:
        return text


def _safe_amount(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text in {"-", "--"}:
        return 0.0
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    text = re.sub(r"[,\s\uc6d0]", "", text)
    text = re.sub(r"[^0-9.\-]", "", text)
    if not text or text in {"-", "."}:
        return 0.0
    try:
        amount = float(text)
    except (TypeError, ValueError):
        return 0.0
    return -amount if negative else amount


def _safe_date(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, _dt.datetime):
        return value.date().isoformat()
    if isinstance(value, _dt.date):
        return value.isoformat()
    text = _safe_text(value)
    if not text:
        return ""
    digits = re.sub(r"\D", "", text)
    if len(digits) == 8:
        return f"{digits[:4]}-{digits[4:6]}-{digits[6:8]}"
    if len(digits) == 6:
        return f"20{digits[:2]}-{digits[2:4]}-{digits[4:6]}"
    return text


def _norm_header(value: Any) -> str:
    text = _repair_mojibake(value).lower()
    return re.sub(r"[\s_()\[\]{}./\\-]+", "", text)


def _unique_headers(worksheet, header_row: int) -> list[str]:
    headers: list[str] = []
    used: dict[str, int] = {}
    for col_idx in range(1, worksheet.max_column + 1):
        raw = _repair_mojibake(worksheet.cell(row=header_row, column=col_idx).value)
        header = raw or get_column_letter(col_idx)
        count = used.get(header, 0) + 1
        used[header] = count
        headers.append(header if count == 1 else f"{header}_{count}")
    return headers


def _find_header_row(worksheet) -> int:
    tokens = (
        "\ud22c\uc790\uc790", "\ucc28\uc785\ucc98", "\uacc4\uc57d", "\ud68c\ucc28",
        "\uc9c0\uae09", "\uc0c1\ud658", "\uc6d0\uae08", "\uc774\uc790",
        "\uc794\uc561", "\uc77c\uc790", "\uae08\uc561",
    )
    best_row = 1
    best_score = -1
    max_scan = min(worksheet.max_row, 30)
    for row_idx in range(1, max_scan + 1):
        values = [
            _norm_header(worksheet.cell(row=row_idx, column=col).value)
            for col in range(1, worksheet.max_column + 1)
        ]
        score = sum(1 for value in values if any(token in value for token in tokens))
        if score > best_score:
            best_score = score
            best_row = row_idx
    return best_row


def _pick(row: dict[str, Any], aliases: tuple[tuple[str, ...], ...]) -> Any:
    normalized = {key: _norm_header(key) for key in row.keys()}
    for alias in aliases:
        targets = tuple(_norm_header(token) for token in alias if _norm_header(token))
        for key, header in normalized.items():
            if targets and all(target in header for target in targets):
                value = row.get(key)
                if value is not None and _safe_text(value) != "":
                    return value
    return None


def _cell(worksheet, row_idx: int, letters: tuple[str, ...]) -> Any:
    for letter in letters:
        col_idx = openpyxl.utils.column_index_from_string(letter)
        if col_idx > worksheet.max_column:
            continue
        value = worksheet.cell(row=row_idx, column=col_idx).value
        if value is not None and _safe_text(value) != "":
            return value
    return None


def _round_label(value: Any) -> str:
    text = _repair_mojibake(value)
    match = re.search(r"(\d+)", text)
    return f"{match.group(1)}\ucc28" if match else text


def _cell_value(value: Any) -> Any:
    if isinstance(value, _dt.datetime):
        return value.date().isoformat()
    if isinstance(value, _dt.date):
        return value.isoformat()
    if isinstance(value, str):
        return _repair_mojibake(value)
    return value


def _is_summary_row(raw: dict[str, Any]) -> bool:
    text = " ".join(_safe_text(value) for value in raw.values())
    return any(token in text for token in ("\ud569\uacc4", "\ucd1d\uacc4", "\uc18c\uacc4"))


def parse_fms_payment_list(filepath: str, period: str = "") -> dict:
    workbook = openpyxl.load_workbook(filepath, data_only=True)
    worksheet = workbook.active
    header_row = _find_header_row(worksheet)
    headers = _unique_headers(worksheet, header_row)

    rows: list[dict[str, Any]] = []
    totals = {
        "borrow_amount": 0.0,
        "payment_amount": 0.0,
        "principal_amount": 0.0,
        "interest_amount": 0.0,
        "balance_amount": 0.0,
    }

    for row_idx in range(header_row + 1, worksheet.max_row + 1):
        raw: dict[str, Any] = {}
        has_value = False
        for col_idx, header in enumerate(headers, start=1):
            value = worksheet.cell(row=row_idx, column=col_idx).value
            if value is not None and _safe_text(value) != "":
                has_value = True
            raw[header] = _cell_value(value)
            raw[f"__col_{get_column_letter(col_idx)}"] = _cell_value(value)
        if not has_value or _is_summary_row(raw):
            continue

        investor_name = _repair_mojibake(
            _cell(worksheet, row_idx, ("A",)) or _pick(raw, (
                ("\uc5c5\uccb4\uba85",), ("\ud22c\uc790\uc790\uba85",), ("\ucc28\uc785\ucc98",), ("\ud22c\uc790\uc790",),
            ))
        )
        contract_name = _repair_mojibake(
            _cell(worksheet, row_idx, ("B",)) or _pick(raw, (
                ("\uacc4\uc57d\uba85",), ("\ucc28\uc785", "\ud68c\ucc28"), ("\ud68c\ucc28",),
            ))
        )
        round_value = contract_name or _pick(raw, (("\ud68c\ucc28",), ("\uacc4\uc57d\uba85",))) or _cell(worksheet, row_idx, ("B", "C", "A"))
        payment_date = _safe_date(
            _cell(worksheet, row_idx, ("C",)) or _pick(raw, (
                ("\uc9c0\uae09", "\uc77c\uc790"), ("\uc9c0\uae09\uc77c",),
                ("\uc0c1\ud658", "\uc77c\uc790"), ("\uc0c1\ud658\uc77c",), ("\uc77c\uc790",),
            ))
        )
        interest_amount = _safe_amount(
            _cell(worksheet, row_idx, ("F",)) or _pick(raw, (
                ("\uc9c0\uae09\uc774\uc790",), ("\uc608\uc815\uc774\uc790",), ("\uc774\uc790",),
            ))
        )
        principal_amount = _safe_amount(
            _cell(worksheet, row_idx, ("G",)) or _pick(raw, (
                ("\uc9c0\uae09\uc6d0\uae08",), ("\uc608\uc815\uc6d0\uae08",), ("\uc6d0\uae08",),
            ))
        )
        borrow_amount = _safe_amount(
            _cell(worksheet, row_idx, ("O",)) or _pick(raw, (
                ("\ucc28\uc785\uae08\uc561",), ("\ucc28\uc785", "\uae08\uc561"),
            ))
        )
        contract_date = _safe_date(
            _cell(worksheet, row_idx, ("P",)) or _pick(raw, (
                ("\uacc4\uc57d\uc77c",), ("\uacc4\uc57d\uc77c\uc790",),
            ))
        )
        maturity_date = _safe_date(
            _cell(worksheet, row_idx, ("Q",)) or _pick(raw, (
                ("\ub9cc\ub8cc\uc77c",), ("\ub9cc\uae30\uc77c",), ("\ub9cc\uae30\uc77c\uc790",),
            ))
        )
        agreement_day = _safe_text(
            _cell(worksheet, row_idx, ("S",)) or _pick(raw, (
                ("\uc57d\uc815\uc77c",),
            ))
        )
        category = _repair_mojibake(
            _cell(worksheet, row_idx, ("T",)) or _pick(raw, (
                ("\ucc28\uc785\ucc98\uad6c\ubd84",), ("\uc720\ud615",), ("\uad6c\ubd84",),
            ))
        )
        balance_amount = _safe_amount(_pick(raw, (
            ("\uc608\uc815\uc794\uc561",), ("\uc794\uc561",), ("\ub0a8\uc740\uc794\uc561",),
        )) or _cell(worksheet, row_idx, ("M", "N")))
        payment_amount = principal_amount + interest_amount
        if payment_amount == 0:
            payment_amount = _safe_amount(_pick(raw, (
                ("\uc9c0\uae09\uc608\uc815\uae08\uc561",), ("\uc9c0\uae09", "\uae08\uc561"),
                ("\uc9c0\uae09\uc561",), ("\ud569\uacc4",),
            )))

        if not investor_name and not contract_name and not payment_date and payment_amount == 0:
            continue

        payment_period = payment_date[:7] if re.match(r"^\d{4}-\d{2}", payment_date or "") else period

        row = {
            "period": payment_period,
            "investor_name": investor_name,
            "borrower": investor_name,
            "contract_name": contract_name,
            "round_label": _round_label(round_value or contract_name),
            "payment_date": payment_date,
            "borrow_amount": borrow_amount,
            "payment_amount": payment_amount,
            "principal_amount": principal_amount,
            "interest_amount": interest_amount,
            "balance_amount": balance_amount,
            "contract_date": contract_date,
            "maturity_date": maturity_date,
            "agreement_day": agreement_day,
            "category": category,
            "raw": raw,
        }
        rows.append(row)
        totals["borrow_amount"] += borrow_amount
        totals["payment_amount"] += payment_amount
        totals["principal_amount"] += principal_amount
        totals["interest_amount"] += interest_amount
        totals["balance_amount"] += balance_amount

    periods = sorted({
        str(row.get("period") or "").strip()
        for row in rows
        if isinstance(row, dict) and str(row.get("period") or "").strip()
    })

    return {
        "period": period,
        "periods": periods or ([period] if period else []),
        "rows": rows,
        "totals": totals,
        "total_payment_amount": totals["payment_amount"],
        "source": "fms_payment",
        "header_row": header_row,
        "headers": headers,
        "total_borrow_amount": totals["borrow_amount"],
    }
