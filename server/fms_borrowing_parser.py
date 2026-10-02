from __future__ import annotations

import re

import openpyxl
from openpyxl.utils import column_index_from_string


def _safe_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _repair_mojibake(value) -> str:
    text = _safe_text(value)
    if not text:
        return ""
    try:
        repaired = text.encode("latin-1").decode("utf-8")
        return repaired if repaired else text
    except UnicodeError:
        return text


def _safe_amount(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except (TypeError, ValueError):
        return 0.0


def _round_label(value) -> str:
    text = _repair_mojibake(value)
    match = re.search(r"(\d+)\s*", text)
    return f"{match.group(1)}\ucc28" if match else text


def parse_fms_borrowing_list(filepath: str, period: str = "") -> dict:
    workbook = openpyxl.load_workbook(filepath, data_only=True)
    worksheet = workbook.active
    cols = {
        "investor_name": column_index_from_string("B"),
        "contract_name": column_index_from_string("C"),
        "borrow_balance": column_index_from_string("M"),
    }

    rows: list[dict] = []
    total_balance = 0.0
    for row_idx in range(2, worksheet.max_row + 1):
        investor_name = _repair_mojibake(worksheet.cell(row=row_idx, column=cols["investor_name"]).value)
        contract_name = _repair_mojibake(worksheet.cell(row=row_idx, column=cols["contract_name"]).value)
        borrow_balance = _safe_amount(worksheet.cell(row=row_idx, column=cols["borrow_balance"]).value)
        if not investor_name and not contract_name and borrow_balance == 0:
            continue
        if not investor_name or not contract_name:
            continue
        total_balance += borrow_balance
        rows.append({
            "period": period,
            "investor_name": investor_name,
            "borrower": investor_name,
            "contract_name": contract_name,
            "round_label": _round_label(contract_name),
            "borrow_balance": borrow_balance,
        })

    return {
        "period": period,
        "periods": [period] if period else [],
        "rows": rows,
        "total_balance": total_balance,
        "source": "fms_borrowing",
    }
