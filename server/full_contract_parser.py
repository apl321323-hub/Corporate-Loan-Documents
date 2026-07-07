"""
Full contract list parser for daily close reports.

Columns used:
  A  : contract number
  C  : product name
  L  : advertising media / agent channel
  V  : loan balance, won
  AD : contract date
  AH : overdue days
  CK : collateral kind, optional
  EK : NICE score, optional
  EO : K score, optional

The parsed shape intentionally mirrors the settlement asset-quality subset used
by the daily close screen: section1, section3, section4 and loan_loss_rows.
"""

from __future__ import annotations

from collections import defaultdict
import re
from datetime import date, datetime

import openpyxl
from openpyxl.utils import column_index_from_string


def _safe_amt(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except (TypeError, ValueError):
        return 0.0


def _safe_days(value) -> int:
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return max(0, int(value))
    match = re.search(r"-?\d+", str(value).replace(",", ""))
    if not match:
        return 0
    try:
        return max(0, int(match.group(0)))
    except ValueError:
        return 0


def _safe_int(value):
    if value is None:
        return None
    try:
        return int(float(str(value).strip().replace(",", "")))
    except (TypeError, ValueError):
        return None


def _contract_no(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def _date_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    match = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", text)
    if match:
        return f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    match = re.search(r"(?<!\d)(\d{2})\D+(\d{1,2})\D+(\d{1,2})(?!\d)", text)
    if match:
        return f"20{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    return ""


def _normalize_period(value: str | None) -> str:
    text = str(value or "").strip()
    match = re.search(r"(\d{4})\D+(\d{1,2})", text)
    if match:
        return f"{match.group(1)}-{int(match.group(2)):02d}"
    match = re.search(r"(?<!\d)(\d{2})\D+(\d{1,2})(?!\d)", text)
    if match:
        return f"20{match.group(1)}-{int(match.group(2)):02d}"
    return text


def _group_products(group: dict) -> list[str]:
    items = group.get("products")
    if items is None:
        items = group.get("items")
    return [str(item).strip() for item in (items or []) if str(item).strip()]


def _product_group_map(product_groups: list[dict] | None) -> tuple[dict[str, str], list[str]]:
    product_to_group: dict[str, str] = {}
    group_names: list[str] = []
    for group in product_groups or []:
        name = str(group.get("name", "")).strip()
        if not name:
            continue
        group_names.append(name)
        for product in _group_products(group):
            product_to_group[product] = name
    return product_to_group, group_names


def _sheet_for_contracts(workbook):
    for sheet_name in workbook.sheetnames:
        if "계약" in sheet_name or "진행" in sheet_name:
            return workbook[sheet_name]
    return workbook.active


def _build_from_rows(rows: list[dict], period: str = "", group_names: list[str] | None = None) -> dict:
    total_won = 0.0
    product_totals: dict[str, float] = defaultdict(float)
    group_totals: dict[str, float] = defaultdict(float)
    products: set[str] = set()
    groups_seen: set[str] = set()

    for row in rows:
        product = str(row.get("product", "")).strip()
        group = str(row.get("group", "")).strip()
        balance = _safe_amt(row.get("balance"))
        total_won += balance
        if product:
            products.add(product)
            product_totals[product] += balance
        if group:
            groups_seen.add(group)
            group_totals[group] += balance

    ordered_groups = list(group_names or [])
    for group in sorted(groups_seen):
        if group not in ordered_groups:
            ordered_groups.append(group)

    section3 = {
        group: {"융잔합계": group_totals.get(group, 0.0) / 1_000_000}
        for group in ordered_groups
    }
    section4 = {
        product: {"융잔합계": product_totals[product] / 1_000_000}
        for product in sorted(products)
    }

    periods = [period] if period else sorted({str(row.get("period", "")).strip() for row in rows if row.get("period")})
    return {
        "period": period,
        "periods": periods,
        "products": sorted(products),
        "section1": {"융잔합계": total_won / 1_000_000},
        "section3": section3,
        "section4": section4,
        "loan_loss_rows": rows,
        "rows": len(rows),
        "source": "full_contract",
    }


def parse_full_contract_list(filepath: str, product_groups: list[dict] | None = None, period: str = "") -> dict:
    period_key = _normalize_period(period)
    workbook = openpyxl.load_workbook(filepath, data_only=True)
    worksheet = _sheet_for_contracts(workbook)
    product_to_group, group_names = _product_group_map(product_groups)
    cols = {
        "contract_no": column_index_from_string("A"),
        "product": column_index_from_string("C"),
        "channel": column_index_from_string("L"),
        "balance": column_index_from_string("V"),
        "contract_date": column_index_from_string("AD"),
        "overdue_days": column_index_from_string("AH"),
        "collateral_kind": column_index_from_string("CK"),
        "nice_score": column_index_from_string("EK"),
        "k_score": column_index_from_string("EO"),
    }

    rows: list[dict] = []
    for row_idx in range(2, worksheet.max_row + 1):
        contract_no = _contract_no(worksheet.cell(row=row_idx, column=cols["contract_no"]).value)
        product = str(worksheet.cell(row=row_idx, column=cols["product"]).value or "").strip()
        balance = _safe_amt(worksheet.cell(row=row_idx, column=cols["balance"]).value)
        if not product and balance == 0:
            continue
        if not product:
            continue
        group = product_to_group.get(product, "")
        channel = str(worksheet.cell(row=row_idx, column=cols["channel"]).value or "").strip()
        contract_date = _date_str(worksheet.cell(row=row_idx, column=cols["contract_date"]).value)
        overdue_days = _safe_days(worksheet.cell(row=row_idx, column=cols["overdue_days"]).value)
        collateral_kind = str(worksheet.cell(row=row_idx, column=cols["collateral_kind"]).value or "").strip()
        nice_score = _safe_int(worksheet.cell(row=row_idx, column=cols["nice_score"]).value)
        k_score = _safe_int(worksheet.cell(row=row_idx, column=cols["k_score"]).value)
        rows.append({
            "period": period_key,
            "contract_no": contract_no,
            "contract_date": contract_date,
            "contract_period": contract_date[:7] if contract_date else "",
            "product": product,
            "group": group,
            "channel": channel,
            "channel_raw": channel,
            "days": overdue_days,
            "balance": balance,
            "balance_million": balance / 1_000_000,
            "nice_score": nice_score,
            "k_score": k_score,
            "collateral_kind": collateral_kind,
        })

    parsed = _build_from_rows(rows, period_key, group_names)
    parsed["source_file"] = filepath
    parsed["column_map"] = cols
    return parsed


def build_full_contract_period(data: dict, period: str) -> dict:
    period_key = _normalize_period(period)
    rows = [
        dict(row)
        for row in (data.get("loan_loss_rows") or [])
        if not period_key or str(row.get("period", "")).strip() == period_key
    ]
    group_names = list((data.get("section3") or {}).keys())
    return _build_from_rows(rows, period_key, group_names)
